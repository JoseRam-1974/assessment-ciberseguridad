import streamlit as st
import pandas as pd
from docx import Document
from fpdf import FPDF
import re
import os

# --- 1. CONFIGURACIÓN Y URL DE LA APP ---
st.set_page_config(page_title="SecureSoft GTD | Assessment Digital", page_icon="🛡️", layout="wide")

# Obtener URL para compartir automáticamente
try:
    current_url = f"https://{st.context.headers.get('host')}"
except:
    current_url = "assessment-ciberseguridad.streamlit.app"

st.markdown(f"""
    <style>
    .stApp {{ background-color: #0b111b; color: #ffffff; }}
    .share-link {{
        background-color: #161f2d;
        padding: 10px;
        border-radius: 5px;
        border: 1px dashed #00adef;
        margin-bottom: 20px;
        font-family: monospace;
    }}
    /* ... (tus estilos anteriores) ... */
    </style>
    """, unsafe_allow_html=True)

# --- 2. FUNCIONES TÉCNICAS ---
def leer_word(ruta):
    try:
        doc = Document(ruta)
        datos = []
        for tabla in doc.tables:
            for fila in tabla.rows:
                celdas = [c.text.strip() for c in fila.cells]
                if len(celdas) >= 2: datos.append([celdas[0], celdas[1]])
        return pd.DataFrame(datos[1:], columns=["Clave", "Contenido"])
    except: return pd.DataFrame()

def clean_pdf(txt):
    if not txt: return ""
    rep = {"á":"a","é":"e","í":"i","ó":"o","ú":"u","ñ":"n","Á":"A","É":"E","Í":"I","Ó":"O","Ú":"U","Ñ":"N","¿":"","¡":"","–":"-"}
    t = str(txt)
    for a, b in rep.items(): t = t.replace(a, b)
    return t.encode('latin-1', 'ignore').decode('latin-1')

class PDF(FPDF):
    def header(self):
        logo = 'Logotipo-SECURESOFT-GTD-Color-Fondo-Transparente.png'
        if os.path.exists(logo):
            self.image(logo, 15, 10, 40)
        self.set_font('Arial', 'B', 10)
        self.set_text_color(0, 85, 165)
        self.set_xy(110, 10)
        self.cell(85, 10, 'ASSESSMENT DIGITAL ESTADO DE CIBERSEGURIDAD', 0, 1, 'R')
        self.ln(20)

# --- 3. INICIO Y REGISTRO ---
if 'etapa' not in st.session_state:
    st.session_state.update({'etapa': 'registro', 'paso': 0, 'respuestas_texto': [], 'preguntas_texto': [], 'datos_usuario': {}})

if st.session_state.etapa == 'registro':
    # BLOQUE PARA COMPARTIR
    st.markdown(f"""
    <div class="share-link">
        🔗 <b>Enlace para compartir:</b> {current_url}
    </div>
    """, unsafe_allow_html=True)
    
    logo_path = 'Logotipo-SECURESOFT-GTD-Color-Fondo-Transparente.png'
    if os.path.exists(logo_path): st.image(logo_path, width=350)
    st.markdown('<p class="cyber-main-title">Assessment Digital Estado de Ciberseguridad</p>', unsafe_allow_html=True)
    
    col1, col2 = st.columns(2)
    with col1:
        nom = st.text_input("Nombre Completo")
        car = st.text_input("Cargo")
        emp = st.text_input("Empresa")
    with col2:
        ema = st.text_input("Email Corporativo")
        tel = st.text_input("Teléfono de Contacto")
        ind = st.text_input("Industria")
    
    if st.button("INICIAR ASSESSMENT"):
        if all([nom, car, emp, ema, tel]):
            st.session_state.datos_usuario = {"Nombre": nom, "Cargo": car, "Empresa": emp, "Email": ema, "Telefono": tel, "Industria": ind}
            st.session_state.etapa = 'preguntas'
            st.rerun()
        else:
            st.error("Por favor, complete los campos obligatorios.")

# --- 4. PREGUNTAS (Sin cambios necesarios aquí) ---
elif st.session_state.etapa == 'preguntas':
    df_p = leer_word("01. Preguntas.docx")
    if not df_p.empty:
        fila = df_p.iloc[st.session_state.paso]
        st.progress((st.session_state.paso + 1) / len(df_p))
        st.write(f"### {fila['Clave']}")
        opciones = [o.strip() for o in fila['Contenido'].split('\n') if o.strip()]
        es_multiple = any(x in fila['Clave'].lower() for x in ["multiple", "múltiple"])
        if es_multiple:
            ans = st.multiselect("Seleccione:", opciones, key=f"q_{st.session_state.paso}")
        else:
            ans = st.radio("Seleccione:", opciones, index=None, key=f"q_{st.session_state.paso}")
        
        if st.button("CONFIRMAR Y SIGUIENTE", type="primary"):
            if ans:
                st.session_state.preguntas_texto.append(fila['Clave'])
                st.session_state.respuestas_texto.append(", ".join(ans) if isinstance(ans, list) else ans)
                if st.session_state.paso < len(df_p) - 1:
                    st.session_state.paso += 1
                    st.rerun()
                else:
                    st.session_state.etapa = 'resultado'
                    st.rerun()

# --- 5. REPORTE PDF (CORRECCIÓN DEFINITIVA DE RECORTE) ---
elif st.session_state.etapa == 'resultado':
    st.markdown('<p class="cyber-main-title">✅ Evaluación Finalizada</p>', unsafe_allow_html=True)
    
    opcion_contacto = st.radio("¿Desea descargar el informe?", ["Sí, deseo el informe estratégico.", "Deseo consultoría gratuita."], index=0)

    if st.button("GENERAR REPORTE PDF", type="primary"):
        df_rec = leer_word("02. Respuestas.docx")
        pdf = PDF()
        pdf.set_margins(left=20, top=15, right=20) # Márgenes amplios
        pdf.set_auto_page_break(auto=True, margin=20)
        pdf.add_page()
        
        pdf.set_font("Arial", 'B', 14)
        pdf.cell(170, 10, clean_pdf(f"REPORTE: {st.session_state.datos_usuario['Empresa']}"), 0, 1, 'C')
        pdf.ln(5)

        for i in range(len(st.session_state.preguntas_texto)):
            p_text = st.session_state.preguntas_texto[i]
            r_text = st.session_state.respuestas_texto[i]
            
            # Pregunta: Usamos ln=1 para asegurar que el siguiente bloque empiece abajo
            pdf.set_font("Arial", 'B', 10); pdf.set_text_color(50, 50, 50)
            pdf.multi_cell(170, 6, clean_pdf(f"Pregunta {i+1}: {p_text}"), ln=1)
            
            # Resultado: Forzamos la posición X y usamos un ancho controlado
            pdf.ln(1)
            pdf.set_x(20)
            pdf.set_font("Arial", '', 10); pdf.set_text_color(0, 0, 0)
            # LA CLAVE: multi_cell con ancho fijo menor al total y sin ln manual intermedio
            pdf.multi_cell(165, 6, clean_pdf(f"Resultado: {r_text}"), ln=1)
            
            # Recomendaciones
            ids = sorted(list(set(re.findall(r'(\d+\.[a-z])', r_text.lower()))))
            for id_s in ids:
                m_s = df_rec[df_rec['Clave'].str.lower() == id_s]
                if not m_s.empty:
                    txt_s = m_s.iloc[0]['Contenido'].strip()
                    pdf.ln(2)
                    pdf.set_x(25)
                    pdf.set_font("Arial", 'I', 9); pdf.set_text_color(0, 85, 165)
                    pdf.multi_cell(155, 5, clean_pdf(f"Recomendacion ({id_s}): {txt_s}"), border=1)
            pdf.ln(10)

        pdf_bytes = bytes(pdf.output())
        st.download_button(label="📥 DESCARGAR PDF", data=pdf_bytes, file_name="Reporte.pdf", mime="application/pdf")
